
import PyPDF2
import openpyxl
from docx import Document
from textblob import TextBlob
from sklearn.feature_extraction.text import CountVectorizer
from sklearn.decomposition import LatentDirichletAllocation, NMF
from sumy.parsers.plaintext import PlaintextParser
from sumy.nlp.tokenizers import Tokenizer
from sumy.summarizers.lex_rank import LexRankSummarizer
from wordcloud import WordCloud
import matplotlib.pyplot as plt

import pandas as pd



def read_pdf(file):
    text = ""
    try:
        reader = PyPDF2.PdfReader(file)
        for page in reader.pages:
            text += page.extract_text()
    except:
        return ""
    return text


def read_excel(file):
    text = ""
    try:
        workbook = openpyxl.load_workbook(file)
        for sheet in workbook:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell:
                        text += str(cell) + " "
    except:
        return ""
    return text


def read_word(file):
    text = ""
    try:
        doc = Document(file)
        for p in doc.paragraphs:
            text += p.text + "\n"
    except:
        return ""
    return text


def read_csv(file):
    text = ""
    try:
        file.seek(0)   # ⭐ VERY IMPORTANT
        df = pd.read_csv(file)
        for _, row in df.iterrows():
            for cell in row:
                if cell:
                    text += str(cell) + " "
    except:
        return ""
    return text






def analyze_text_stats(text):
    char_count = len(text.replace(" ", "").replace("\n", "").replace("\t", "").replace(".", ""))
    word_count = len(text.split())
    line_count = len([l for l in text.splitlines() if l.strip()])
    sentence_count = len([s for s in text.split('.') if s.strip()])
    return char_count, word_count, line_count, sentence_count



def topic_modeling(text, num_topics=5, model_type="LDA"):
    vectorizer = CountVectorizer(stop_words="english")
    X = vectorizer.fit_transform([text])

    if model_type == "LDA":
        model = LatentDirichletAllocation(n_components=num_topics)
    else:
        model = NMF(n_components=num_topics)

    model.fit(X)
    words = vectorizer.get_feature_names_out()

    topics = []
    for idx, topic in enumerate(model.components_):
        top_words = [words[i] for i in topic.argsort()[-10:]]
        topics.append((f"Topic {idx+1}", ", ".join(top_words)))

    return topics



def analyze_sentiment(text):
    blob = TextBlob(text)
    polarity = blob.sentiment.polarity
    if polarity > 0:
        sentiment = "Positive"
    elif polarity < 0:
        sentiment = "Negative"
    else:
        sentiment = "Neutral"
    return polarity, sentiment



def summarize_text(text, sentences=3):
    parser = PlaintextParser.from_string(text, Tokenizer("english"))
    summarizer = LexRankSummarizer()
    summary = summarizer(parser.document, sentences)
    return " ".join([str(sentence) for sentence in summary])


def generate_wordcloud(processed_text):
    wc = WordCloud(width=800, height=400).generate(processed_text)
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.imshow(wc)
    ax.axis("off")
    return fig

