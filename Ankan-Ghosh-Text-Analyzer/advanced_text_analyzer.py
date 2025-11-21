import streamlit as st
import PyPDF2
import openpyxl
from docx import Document
import re

st.set_page_config(page_title="Text Analyzer", page_icon="📝")

def read_pdf(file):
    text = ""
    try:
        reader = PyPDF2.PdfReader(file)
        for page in reader.pages:
            text += page.extract_text()
    except Exception as e:
        st.error("Cannot read PDF file")
    return text

def read_excel(file):
    text = ""
    try:
        workbook = openpyxl.load_workbook(file)
        for sheet in workbook:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell:
                        text += str(cell) + " "
    except Exception as e:
        st.error("Cannot read Excel file")
    return text

def read_word(file):
    text = ""
    try:
        doc = Document(file)
        for para in doc.paragraphs:
            text += para.text + "\n"
    except Exception as e:
        st.error("Cannot read Word file")
    return text

def clean_text(text):
    """Simple text cleaning"""
    cleaned = re.sub(r'\s+', ' ', text)
    cleaned = cleaned.strip()
    return cleaned

def remove_special_chars(text):
    """Remove special characters and keep only letters, numbers, and spaces"""
    cleaned = re.sub(r'[^a-zA-Z0-9\s]', '', text)
    return cleaned

def analyze_text(text):
    char_count = len(text.replace(" ", "").replace("\n", "").replace("\t", ""))
    
    words = text.split()
    word_count = len(words)
    
    lines = text.split("\n")
    line_count = len(lines)
    
    sentences = text.split('.')
    sentence_count = len([s for s in sentences if s.strip()])
    
    return char_count, word_count, line_count, sentence_count

st.title("📝 AI Text Analyzer")
st.write("Upload a file or type text to analyze")

uploaded_file = st.file_uploader("Choose a file", type=['pdf', 'docx', 'xlsx', 'txt'])

text = ""

if uploaded_file:
    file_type = uploaded_file.name.split('.')[-1]
    
    if file_type == 'pdf':
        text = read_pdf(uploaded_file)
    elif file_type == 'docx':
        text = read_word(uploaded_file)
    elif file_type == 'xlsx':
        text = read_excel(uploaded_file)
    elif file_type == 'txt':
        text = uploaded_file.read().decode('utf-8')
    
    st.success("File uploaded successfully!")

user_text = st.text_area("Type or paste your text here:", value=text, height=250)

st.subheader("Text Cleaning Options")

col1, col2 = st.columns(2)
with col1:
    remove_extra_space = st.checkbox("Remove extra spaces", value=True)
with col2:
    remove_special = st.checkbox("Remove special characters", value=True)

if st.button("Analyze Text"):
    if user_text:
        processed_text = user_text
        
        if remove_extra_space:
            processed_text = clean_text(processed_text)
        
        if remove_special:
            processed_text = remove_special_chars(processed_text)
        
        if remove_extra_space or remove_special:
            with st.expander("View Cleaned Text"):
                st.text_area("Cleaned version:", value=processed_text, height=150, disabled=True)
        
        char_count, word_count, line_count, sentence_count = analyze_text(processed_text)
        
        st.subheader("Results:")
        
        col1, col2 = st.columns(2)
        
        with col1:
            st.metric("Characters (no spaces)", char_count)
            st.metric("Words", word_count)
        
        with col2:
            st.metric("Lines", line_count)
            st.metric("Sentences", sentence_count)
        
    else:
        st.warning("Enter some text or upload a file!")

if st.button("Clear"):
    st.rerun()

st.markdown("---")
st.caption("Text Analyzer with Preprocessing")
